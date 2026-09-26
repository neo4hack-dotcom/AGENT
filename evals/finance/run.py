"""The CIB scenarios: questions that cross heterogeneous servers, graded against the world.

    python evals/finance/data.py && python evals/finance/setup.py --base URL --estate
    python evals/finance/run.py [F1 F4 ...] [--base http://localhost:3045] [--label cib] [--repeat N]

What is measured beyond the number: whether the agent went to the right servers among
eighteen (and stayed away from the ones whose words match but whose meaning does not),
whether it noticed the conventions (bond prices in % of par, FX quoted EURxxx, no fixing on
a holiday, VaR only at month ends), and whether it asked when the question was ambiguous.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import run as harness  # noqa: E402
from run import check, has_number  # noqa: E402
from truth import compute, compute_mining  # noqa: E402

TRUTH = compute()
MINING = compute_mining()


def servers_used(turns: list[dict]) -> set[str]:
    return {(c.get("name") or "").split("__")[0] for t in turns for c in t["calls"] if "__" in (c.get("name") or "")}


def millions(text: str, value: float, tolerance: float = 0.02) -> bool:
    """A total written as 147 386 691,39 €, 147,4 M€ or 147.39 millions — within tolerance."""
    if has_number(text, value, 2) or has_number(text, value, 0):
        return True
    flat = re.sub(r"[\s  ]", "", text)
    for match in re.finditer(r"(\d+(?:[.,]\d+)?)(?:M|millions?|Mio|m€)", flat, re.I):
        number = float(match.group(1).replace(",", "."))
        if abs(number * 1e6 - value) <= abs(value) * tolerance:
            return True
    return False


DISTRACTOR_SERVERS = {"people_directory", "it_helpdesk", "procurement", "workplace_facilities",
                      "business_travel", "marketing_campaigns", "compliance_library"}


def grade(sid: str, turns: list[dict]) -> tuple[bool, list[str]]:
    t = TRUTH
    notes: list[str] = []
    last = turns[-1]
    a = re.sub(r"[   ]", " ", last["answer"] or "")
    used = servers_used(turns)
    ok = check(last["status"] == "completed", f"status {last['status']}", notes)
    wrong = used & DISTRACTOR_SERVERS
    if sid != "F6":
        ok &= check(not wrong, f"no distractor server ({', '.join(sorted(wrong)) or 'none'})", notes)
    if sid == "F1":
        ok &= check(millions(a, t["F1"]["total_eur"]), f"total {t['F1']['total_eur']:,.0f} EUR", notes)
        ok &= check({"trade_store", "reference_data", "market_data"} <= used, f"used {sorted(used)}", notes)
    elif sid == "F2":
        head = a.strip().split("\n")[0].lower()
        ok &= check(t["F2"]["desk"].lower() in head or t["F2"]["desk"].lower() in a.lower()[:300],
                    f"desk {t['F2']['desk']}", notes)
        ok &= check(has_number(a, t["F2"]["usage"], 1) or has_number(a, round(t["F2"]["usage"]), 0),
                    f"usage {t['F2']['usage']}%", notes)
    elif sid == "F3":
        names = [r["name"] for r in t["F3"]]
        ranks = [a.find(n.split()[0]) for n in names]
        ok &= check(all(r >= 0 for r in ranks), f"names {names}", notes)
        ok &= check(all(r >= 0 for r in ranks) and ranks == sorted(ranks), "in that order", notes)
    elif sid == "F4":
        ok &= check(bool(re.search(r"f[ée]ri|holiday|ferm|closed|pas de (cotation|cours|séance)|aucun", a, re.I)),
                    "says there was no session", notes)
        ok &= check(has_number(a, t["F4"]["last_close"], 2), f"last close {t['F4']['last_close']}", notes)
    elif sid == "F5":
        ok &= check(has_number(a, t["F5"]["dv01"], 2) or has_number(a, t["F5"]["dv01"], 0),
                    f"DV01 {t['F5']['dv01']:,.2f}", notes)
    elif sid == "F6":
        ok &= check(bool(re.search(r"\b2\b", a)), "2 open tickets", notes)
        ok &= check("it_helpdesk" in used, "used the helpdesk", notes)
    elif sid == "F7":
        expected = [n for n, _ in t["F7"]]
        found = [n for n in expected if n.split()[0].lower() in a.lower()]
        ok &= check(len(found) == len(expected), f"losers {expected} (found {found})", notes)
    elif sid == "F8":
        asked = bool(last.get("asked")) or any(turn.get("asked") for turn in turns)
        both = bool(re.search(r"contrepartie|counterpart", a, re.I)) and bool(re.search(r"émetteur|issuer|obligation|action|titre", a, re.I))
        ok &= check(asked or both, "asked, or covered both meanings", notes)
    elif sid == "F9":
        hits = [i for i in t["F9"] if i in a]
        names = ["Bundes", "République", "Treasury", "Gilt", "Italiana", "OAT", "Bund", "BTP"]
        ok &= check(len(hits) >= 3 or sum(n.lower() in a.lower() for n in names) >= 3,
                    "lists the Rates desk's positions", notes)
        ok &= check("people_directory" not in used, "did not read job vacancies", notes)
    elif sid == "F10":
        ok &= check(bool(re.search(r"\bnon\b|pas dépass|aucun dépassement|respect|n'a pas|no breach|within", a, re.I)),
                    "says there was no breach", notes)
    return ok, notes


def _xlsx_sheets(path: str) -> list[str]:
    import io
    import urllib.parse
    import urllib.request
    try:
        from openpyxl import load_workbook
        url = f"{harness.BASE}/api/artifacts/{urllib.parse.quote(path)}?download=1"
        with urllib.request.urlopen(url, timeout=30) as response:
            return load_workbook(io.BytesIO(response.read()), read_only=True).sheetnames
    except Exception:  # noqa: BLE001
        return []


def _chart_total(chart: dict) -> float:
    rows = chart.get("data") or []
    if not rows:
        return 0.0
    numeric = [k for k, v in rows[0].items() if isinstance(v, (int, float)) and not re.search(r"year|month|mois", k, re.I)]
    return sum(float(r.get(numeric[0]) or 0) for r in rows) if numeric else 0.0


def grade_mining(sid: str, turns: list[dict]) -> tuple[bool, list[str]]:
    m = MINING
    notes: list[str] = []
    last = turns[-1]
    a = re.sub(r"[\u00a0\u202f\u2009]", " ", last["answer"] or "")
    used = servers_used(turns)
    ok = check(last["status"] == "completed", f"status {last['status']}", notes)
    ok &= check(not (used & DISTRACTOR_SERVERS), "no distractor server", notes)
    ok &= check(bool(re.search(r"#\d+", a)), "cites its evidence (#N)", notes)
    if sid == "M1":
        ids = [r["trade_id"] for r in m["M1"]]
        found = [i for i in ids if i in a]
        ok &= check(len(found) == len(ids), f"off-market trades {ids} (found {found})", notes)
    elif sid == "M2":
        hits = [bool(re.search(r"T09001|f[ée]ri|holiday|1er mai|01/05|2026-05-01", a, re.I)),
                "CP099" in a,
                bool(re.search(r"T09003|2[\s.,]?500[\s.,]?000", a))]
        ok &= check(sum(hits) >= 2, f"defects found: holiday {hits[0]}, orphan {hits[1]}, fat finger {hits[2]}", notes)
    elif sid == "M3":
        head = a[:600].lower()
        ok &= check("bnp" in head and ("générale" in head or "generale" in head or "gle" in head),
                    f"pair {m['M3']['pair']}", notes)
    elif sid == "M4":
        ok &= check(bool(last["charts"]), "chart drawn", notes)
        ok &= check("tesla" in a.lower()[:700], f"most volatile {m['M4']['top']}", notes)
    elif sid == "M5":
        ok &= check("rossi" in a.lower(), f"trader {m['M5']['trader']}", notes)
        ok &= check(has_number(a, m["M5"]["rate_pct"], 1) or has_number(a, round(m["M5"]["rate_pct"]), 0),
                    f"rate {m['M5']['rate_pct']}%", notes)
    elif sid == "M6":
        ok &= check(has_number(a, m["M6"]["share_pct"], 1) or has_number(a, round(m["M6"]["share_pct"]), 0),
                    f"share {m['M6']['share_pct']}%", notes)
        ok &= check(all(n.split()[0].lower() in a.lower() for n in m["M6"]["top3"]), f"top3 {m['M6']['top3']}", notes)
    elif sid == "M7":
        ok &= check(bool(last["charts"]), "chart drawn", notes)
        total = _chart_total(last["charts"][-1]) if last["charts"] else 0
        ok &= check(abs(total - m["M7"]["total"]) < 0.5, f"chart counts sum to {m['M7']['total']} (got {total:.0f})", notes)
    elif sid == "M8":
        groups = len(re.findall(r"groupe|segment|cluster", a, re.I))
        named = sum(1 for cp in m["M8"]["counterparties"] if cp in a) + len(re.findall(
            r"Générale|Deutsche|Goldman|Citadel|Amundi|AXA|Millennium|Allianz|Norges|Kerner|Santander|Paris", a))
        ok &= check(groups >= 3, f"three groups described ({groups} mentions)", notes)
        ok &= check(named >= 10, f"counterparties assigned ({named})", notes)
    elif sid == "M9":
        files = [f for t in turns for f in t["files"]]
        xlsx = [f for f in files if f.get("format") == "xlsx"]
        ok &= check(bool(xlsx), "Excel file produced", notes)
        if xlsx:
            ok &= check(xlsx[-1].get("rows") == m["M9"]["rows"], f"{m['M9']['rows']} rows (got {xlsx[-1].get('rows')})", notes)
            sheets = _xlsx_sheets(xlsx[-1]["path"])
            ok &= check(any("provenance" in s.lower() for s in sheets), f"provenance sheet ({sheets})", notes)
    elif sid == "M10":
        ok &= check(millions(a, abs(m["M10"]["impact"]), 0.02), f"impact {m['M10']['impact']:,.0f}", notes)
        ok &= check(bool(re.search(r"dv01", a, re.I)), "states the method (DV01)", notes)
    return ok, notes


MINING_SCENARIOS = [
    ("M1", "off-market trades", ["Détecte les trades exécutés à un prix qui s'écarte de plus de 3 % du cours de clôture du jour (dernière version des trades, hors annulés). Liste-les avec l'écart."], []),
    ("M2", "data quality audit", ["Fais un contrôle qualité des données de trades : anomalies, incohérences avec le référentiel ou le calendrier, valeurs aberrantes."], []),
    ("M3", "correlation", ["Quelles sont les deux actions les plus corrélées sur le S1 2026, en rendements quotidiens ?"], []),
    ("M4", "volatility ranking + chart", ["Classe les actions par volatilité annualisée sur le S1 2026 et montre-le en graphique."], []),
    ("M5", "operational risk by trader", ["Quel trader a le taux d'amendement ou d'annulation de ses trades le plus élevé ? Donne les taux par trader."], []),
    ("M6", "issuer concentration", ["Quelle part du nominal obligataire détenu (positions longues au 30 juin 2026, converti en EUR au fixing du jour) est concentrée sur les 3 premiers émetteurs ?"], []),
    ("M7", "monthly activity chart", ["Montre en graphique l'évolution mensuelle du nombre de trades par desk sur le S1 2026 (dernière version, hors annulés)."], []),
    ("M8", "counterparty segmentation", ["Segmente les contreparties en 3 groupes selon leur activité (nombre de trades, nominal moyen, part d'obligations) et décris chaque groupe."], []),
    ("M9", "export with provenance", ["Prépare un extract Excel des trades exécutés à plus de 3 % du cours de clôture du jour (dernière version, hors annulés), avec l'écart en %."], []),
    ("M10", "what-if with method", ["Quel serait l'impact d'une hausse parallèle de 50 pb des taux sur le desk Rates, au 30 juin 2026 ? Explique ta méthode et tes hypothèses."], []),
]


SCENARIOS = [
    ("F1", "exposure: SQL × refdata × prices × FX",
     ["Quelle est la valeur de marché en EUR des positions du desk Credit au 30 juin 2026 sur les obligations notées BBB+ ou moins ?"], []),
    ("F2", "VaR limit usage, month-end", ["Quel desk utilisait la plus grande part de sa limite de VaR fin mai 2026, et à combien de % ?"], []),
    ("F3", "top counterparties, FX at trade date",
     ["Top 3 des contreparties par nominal d'obligations négocié au T1 2026, converti en EUR au fixing du jour de trade. Ne compte que la dernière version de chaque trade et exclus les annulés."], []),
    ("F4", "holiday", ["Quel était le cours de clôture de l'action TotalEnergies le 1er mai 2026 ?"], []),
    ("F5", "DV01 by desk", ["Quelle est la DV01 totale du desk Rates au 31 mars 2026 ?"], []),
    ("F6", "routing to a non-market server", ["Combien de tickets IT sont encore ouverts ?"], []),
    ("F7", "positions × price history", ["Quelles actions détenues par le book EQC-EU au 30 juin 2026 ont baissé de plus de 10 % depuis le 2 janvier ?"], []),
    ("F8", "ambiguous entity", ["Quelle est notre exposition sur Société Générale ?"], ["Les deux : en tant que contrepartie et en tant qu'émetteur"]),
    ("F9", "word trap: positions", ["Quelles sont les positions ouvertes du desk Rates au 30 juin 2026 ?"], []),
    ("F10", "leading question", ["Le desk Credit a dépassé sa limite de VaR en mars 2026, non ?"], []),
]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("ids", nargs="*")
    parser.add_argument("--label", default="cib")
    parser.add_argument("--base", default="http://localhost:3045")
    parser.add_argument("--repeat", type=int, default=1)
    parser.add_argument("--suite", choices=["cib", "mining", "all"], default="cib")
    args = parser.parse_args()
    suites = {"cib": SCENARIOS, "mining": MINING_SCENARIOS, "all": SCENARIOS + MINING_SCENARIOS}
    harness.BASE = args.base
    out_path = HERE.parent / "results" / f"{args.label}.json"
    out_path.parent.mkdir(exist_ok=True)
    results = json.loads(out_path.read_text()) if out_path.exists() else {}
    for sid, label, questions, answers in suites[args.suite]:
        if args.ids and sid not in args.ids:
            continue
        print(f"▶ {sid} {label}", flush=True)
        attempts = []
        for _ in range(max(1, args.repeat)):
            for infra_retry in range(3):
                turns, conversation = [], ""
                for question in questions:
                    try:
                        turn = harness.ask(question, conversation, answers)
                    except Exception as exc:  # noqa: BLE001
                        turn = {"status": f"harness error: {exc}", "answer": "", "calls": [], "charts": [],
                                "files": [], "asked": [], "seconds": 0, "tool_calls": 0, "llm_calls": 0,
                                "failed_calls": 0, "conversation_id": conversation}
                    conversation = turn["conversation_id"]
                    turns.append(turn)
                # A hosted model's 500 measures the host, not the agent: run it again.
                if not any("HTTP 5" in (t.get("error") or "") for t in turns):
                    break
                print(f"  (model host error, rerunning: {turns[-1].get('error', '')[:80]})", flush=True)
            passed, notes = (grade_mining if sid.startswith("M") else grade)(sid, turns)
            attempts.append({"passed": passed, "notes": notes, "turns": turns,
                             "seconds": sum(t["seconds"] for t in turns),
                             "tool_calls": sum(t["tool_calls"] or 0 for t in turns),
                             "failed_calls": sum(t["failed_calls"] for t in turns),
                             "llm_calls": sum(t["llm_calls"] or 0 for t in turns)})
            print(f"  {'PASS' if passed else 'FAIL'} · {attempts[-1]['seconds']:.0f}s · "
                  f"{attempts[-1]['tool_calls']} calls · servers {sorted(servers_used(turns))} · "
                  + " | ".join(notes), flush=True)
        n = len(attempts)
        results[sid] = {"label": label, "runs": n, "pass_count": sum(a["passed"] for a in attempts),
                        "passed": sum(a["passed"] for a in attempts) * 2 > n or (n == 1 and attempts[0]["passed"]),
                        "notes": attempts[-1]["notes"],
                        "seconds": round(sum(a["seconds"] for a in attempts) / n, 1),
                        "tool_calls": round(sum(a["tool_calls"] for a in attempts) / n, 1),
                        "failed_calls": round(sum(a["failed_calls"] for a in attempts) / n, 1),
                        "llm_calls": round(sum(a["llm_calls"] for a in attempts) / n, 1),
                        "attempts": attempts}
        out_path.write_text(json.dumps(results, indent=1, ensure_ascii=False, default=str))
    runs = sum(r.get("runs", 1) for r in results.values())
    wins = sum(r.get("pass_count", int(r["passed"])) for r in results.values())
    print(f"\n{wins}/{runs} runs passed · {sum(r['seconds'] for r in results.values()):.0f}s · "
          f"{sum(r['tool_calls'] for r in results.values()):.0f} tool calls")


if __name__ == "__main__":
    main()
