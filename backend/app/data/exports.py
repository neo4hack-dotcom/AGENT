"""Files the reader takes away: extracts people open in Excel, not dumps they have to clean.

An extract is judged in the first ten seconds after it is opened. So the header is frozen
and filterable, numbers are numbers with thousands separators, ISO dates are real dates
Excel can sort and pivot, columns are wide enough to read, and the workbook says what it
is in its properties. CSV is written with a byte-order mark because Excel on Windows still
reads UTF-8 without one as Latin-1 and turns "Société" into mojibake.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import re
from pathlib import Path

_ISO_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_TIME = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?")


class ExportError(ValueError):
    pass


# Names Windows keeps for devices, extension or not: `nul.csv` is the null device there.
_RESERVED = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(10)), *(f"lpt{i}" for i in range(10))}


def safe_name(name: str, extension: str) -> str:
    stem = re.sub(r"[^\w\-. ]+", "", (name or "export").strip(), flags=re.UNICODE)
    stem = re.sub(r"\s+", "-", stem).strip(".-") or "export"
    stem = re.sub(rf"\.{extension}$", "", stem, flags=re.IGNORECASE)[:80].rstrip(". ")
    if stem.split(".")[0].lower() in _RESERVED:
        stem = f"{stem}-export"
    return f"{stem or 'export'}.{extension}"


def unique_path(folder: Path, name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / name
    stem, suffix, n = target.stem, target.suffix, 2
    while target.exists():
        target = folder / f"{stem}-{n}{suffix}"
        n += 1
    return target


def _columns(rows: list[dict]) -> list[str]:
    seen: dict[str, None] = {}
    for row in rows:
        for key in row:
            seen.setdefault(str(key), None)
    return list(seen)


def write_csv(rows: list[dict], path: Path) -> None:
    columns = _columns(rows)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: ("" if row.get(k) is None else row.get(k)) for k in columns})


def write_json(rows: list[dict], path: Path) -> None:
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


def _provenance_sheet(book, provenance: dict) -> None:
    """A last sheet that says where every other sheet came from."""
    from openpyxl.styles import Alignment, Font
    sheet = book.create_sheet("Provenance")
    bold = Font(bold=True)
    rows = [("Question", provenance.get("question") or ""),
            ("Generated at", provenance.get("generated_at") or ""),
            ("Model", provenance.get("model") or ""), ("", "")]
    for key, value in rows:
        sheet.append([key, value])
        sheet.cell(sheet.max_row, 1).font = bold
    sheet.append(["Sheet", "Step", "Ref", "Source", "Tool", "Operation (query, code or arguments)",
                  "Rows", "Uses", "Retrieved at", "Result fingerprint (sha256)", "Audit log entry"])
    for cell in sheet[sheet.max_row]:
        cell.font = bold
    for table in provenance.get("tables") or []:
        chain = table.get("chain") or []
        if not chain:
            sheet.append([table.get("sheet"), 1, "", "", "", f"source: {table.get('source')}"])
        for step, node in enumerate(chain, start=1):
            at = node.get("at")
            sheet.append([table.get("sheet"), step, node.get("ref"), node.get("source"), node.get("tool"),
                          node.get("operation"), node.get("rows"), ", ".join(node.get("depends_on") or []),
                          dt.datetime.fromtimestamp(at).strftime("%Y-%m-%d %H:%M:%S") if at else "",
                          node.get("fingerprint"), node.get("audit")])
    for letter, width in zip("ABCDEFGHIJK", (14, 6, 7, 18, 26, 90, 8, 12, 20, 20, 20)):
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=5):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")


def _decimals(rows: list[dict], column: str, floor: int = 0) -> str:
    """The decimals a column needs, as a number-format suffix: none for whole numbers, as
    many as the values carry (up to four) otherwise, at least `floor`."""
    places = 0
    for row in rows:
        value = row.get(column)
        if isinstance(value, float) and not isinstance(value, bool) and not value.is_integer():
            text = f"{round(value, 6):.6f}".rstrip("0")
            places = max(places, len(text.split(".")[1]) if "." in text else 0)
    places = min(4, max(places, floor if places or floor else 0))
    return "." + "0" * places if places else ""


def _signed(fmt: str) -> str:
    """Negatives in red with their minus sign — how a P&L column is read on a desk."""
    return f"{fmt};[Red]-{fmt}"


def write_xlsx(sheets: list[tuple[str, list[dict]]], path: Path, title: str = "",
               provenance: dict | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    header_font = Font(bold=True, color="18181B")
    header_fill = PatternFill("solid", fgColor="EEF2F0")
    used_names: set[str] = set()
    for raw_name, rows in sheets:
        name = re.sub(r"[\[\]:*?/\\]", "", raw_name or "Data")[:31] or "Data"
        base, n = name, 2
        while name in used_names:
            name = f"{base[:28]}-{n}"
            n += 1
        used_names.add(name)
        sheet = book.create_sheet(name)
        columns = _columns(rows)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center")
        widths = [len(str(c)) for c in columns]
        for row in rows:
            values = []
            for i, column in enumerate(columns):
                value = row.get(column)
                if isinstance(value, str) and _ISO_DAY.match(value):
                    try:
                        value = dt.date.fromisoformat(value)
                    except ValueError:
                        pass
                elif isinstance(value, str) and _ISO_TIME.match(value):
                    try:
                        value = dt.datetime.fromisoformat(value.replace("Z", ""))
                    except ValueError:
                        pass
                values.append(value)
                widths[i] = max(widths[i], min(60, len(str(value)) if value is not None else 0))
            sheet.append(values)
        # What each column is — a year, an amount in EUR, a share — read from its values
        # the way charts read them, so 2026 is not written "2,026" and 0.183 reads 18.3 %.
        from app.data.chart_sense import infer
        kinds = infer(rows)
        for i, column in enumerate(columns, start=1):
            letter = get_column_letter(i)
            sheet.column_dimensions[letter].width = min(62, max(9, widths[i - 1] + 2))
            sample = next((r.get(column) for r in rows if r.get(column) is not None), None)
            info = kinds.get(column) or {}
            fmt = None
            if isinstance(sample, bool):
                continue
            if info.get("kind") == "year" or re.search(r"(^|_)id$|code", column, re.I):
                fmt = "0" if isinstance(sample, (int, float)) else None
            elif info.get("kind") == "percent" and isinstance(sample, (int, float)):
                fmt = "0.0%" if info.get("fraction") else '0.0" %"'
            elif info.get("kind") == "currency" and isinstance(sample, (int, float)):
                symbol = {"EUR": ' "€"', "USD": ' "$"', "GBP": ' "£"', "CHF": ' "CHF"', "JPY": ' "¥"'}.get(
                    (info.get("unit") or "").upper(), "")
                fmt = _signed(f"#,##0{_decimals(rows, column, floor=2)}{symbol}")
            elif isinstance(sample, (int, float)):
                # From every value, not the first: [35, 41.5] formatted from 35 showed 41.5 bp
                # as "42" — a display that rounds a spread is a wrong number on screen.
                fmt = _signed(f"#,##0{_decimals(rows, column)}")
            elif isinstance(sample, str) and _ISO_DAY.match(sample):
                fmt = "yyyy-mm-dd"
            if fmt:
                for cell in sheet[letter][1:]:
                    cell.number_format = fmt
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    if provenance:
        _provenance_sheet(book, provenance)
    book.properties.title = title or "AGENT export"
    book.properties.creator = "AGENT"
    book.save(path)


def export(sheets: list[tuple[str, list[dict]]], fmt: str, folder: Path, filename: str,
           title: str = "", provenance: dict | None = None, overwrite: bool = False) -> dict:
    fmt = (fmt or "xlsx").lower().strip(".")
    if fmt not in ("csv", "xlsx", "json"):
        raise ExportError(f"Unknown format '{fmt}'. Use csv, xlsx or json.")
    if not sheets or not any(rows for _, rows in sheets):
        raise ExportError("There are no rows to export.")
    if fmt != "xlsx" and len(sheets) > 1:
        raise ExportError(f"A {fmt.upper()} file holds one table. Use xlsx for several sheets, "
                          f"or export them one at a time.")
    if overwrite:
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / safe_name(filename or title or "export", fmt)
    else:
        path = unique_path(folder, safe_name(filename or title or "export", fmt))
    if fmt == "csv":
        write_csv(sheets[0][1], path)
    elif fmt == "json":
        write_json(sheets[0][1], path)
    else:
        write_xlsx(sheets, path, title, provenance)
    if provenance and fmt in ("csv", "json"):
        # A CSV or JSON file has no room for a second table: its provenance travels beside it.
        import json as _json
        path.with_name(path.stem + ".provenance.json").write_text(
            _json.dumps(provenance, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return {"path": str(path), "name": path.name, "format": fmt, "bytes": path.stat().st_size,
            "rows": sum(len(rows) for _, rows in sheets),
            "sheets": [name for name, _ in sheets] if fmt == "xlsx" else []}
