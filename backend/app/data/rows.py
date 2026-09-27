"""Turn "the result of call #4" into rows, without the rows passing through the model.

Every output the agent produces from data — a chart, an Excel extract, a PDF table — needs
the rows. The model already *has* them, as text in its context, and the tempting design is
to let it pass them along as an argument. That is exactly how a 411-row result became a
20-row file earlier in this project: data retyped through a model comes back shorter,
rounder and quietly invented. So outputs name their data instead — `#4`, a workspace file,
another chart — and this module fetches the real thing.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

MAX_ROWS = 50_000

_REF = re.compile(r"^\s*[\[【]?#(\d{1,3})[\]】]?\s*$")


class SourceError(ValueError):
    """A data reference that does not resolve — said plainly, with what does exist."""


# ----------------------------------------------------------------- parsing text

def rows_from_text(text: str) -> list[dict] | None:
    """Rows from whatever shape a tool answered in, or None if it is not tabular.

    Tolerant on purpose: MCP servers answer with a JSON array of objects, an object
    wrapping one (`rows`, `data`, `records`, `results`), a columns-plus-values pair, CSV,
    or a markdown table. All of them are the same table.
    """
    body = (text or "").strip()
    if not body:
        return None
    if body[:1] in "[{":
        # Whatever notes the runner appended after structured data are not rows.
        from app.agent.context import split_notes
        text = split_notes(body)[0]
    # The runner attaches bracketed notes to some results — before the data (an argument
    # the server ignored) or after it (a whole-table aggregate). Neither is a row.
    paragraphs = [p for p in (text or "").split("\n\n") if not _NOTE.match(p.strip())]
    # Leading spaces are layout in a printed frame: its header is indented to line up with
    # the columns. Everything else parses the stripped text.
    aligned = "\n".join(l.rstrip() for l in "\n\n".join(paragraphs).splitlines() if l.strip())
    body = aligned.strip()
    if body[:1] in "[{":
        try:
            value = json.loads(body)
        except ValueError:
            value = None
        if value is None and len(body) <= 2_000_000:
            # `print(rows)` in run_python: a Python literal — single quotes, True, None.
            # literal_eval builds data and nothing else; no name is looked up, no code runs.
            import ast
            try:
                value = ast.literal_eval(body)
            except (ValueError, SyntaxError, MemoryError, RecursionError):
                value = None
        if value is not None:
            return _rows_from_value(value)
    first = body.splitlines()[0] if body else ""
    if first.lstrip().startswith("|"):
        table = _rows_from_markdown(body)
        if table:
            return table
    printed = _rows_from_printed_frame(aligned)
    if printed:
        return printed
    return _rows_from_csv(body)


_FOOTER = re.compile(r"^\[\d+ rows x \d+ columns\]$")


def _rows_from_printed_frame(text: str) -> list[dict] | None:
    """What `print(df)` produces: a header, then rows that each start with the index.

    The model is told to print JSON for anything it wants to chart, and still prints the
    frame about half the time. Reading the frame is kinder than a failed chart and a second
    run — pandas' own fixed-width reader splits the columns the way pandas printed them.
    """
    lines = [l for l in text.splitlines() if l.strip() and not _FOOTER.match(l.strip())]
    if len(lines) < 2 or len(lines) > MAX_ROWS:
        return None
    body = lines[1:]
    if sum(bool(re.match(r"^\s*\d+\s+\S", l)) for l in body) < 0.8 * len(body):
        return None
    try:
        import pandas as pd
        frame = pd.read_fwf(io.StringIO("\n".join(lines)))
    except Exception:  # noqa: BLE001 - not a printed frame after all
        return None
    if frame.empty or len(frame.columns) < 2:
        return None
    if str(frame.columns[0]).startswith("Unnamed"):
        frame = frame.drop(columns=frame.columns[0])
    records = json.loads(frame.to_json(orient="records"))
    return records or None


_NOTE = re.compile(r"^\[(This aggregates|[A-Za-z0-9_]+__[A-Za-z0-9_]+ does not accept|"
                   r"identical call|result truncated|Several records match|'[a-z_]+' also exists in|"
                   r"This aggregates without)")


def _rows_from_value(value: Any) -> list[dict] | None:
    if isinstance(value, dict) and isinstance(value.get("preview"), list):
        # A dataframe server answers with a description and a preview. The preview is the
        # table only when it is the whole table; a partial one charted as if it were
        # complete is a silently wrong chart.
        total = value.get("rows") if isinstance(value.get("rows"), int) else None
        preview = value["preview"]
        if total is not None and total > len(preview):
            raise SourceError(f"That result shows only the first {len(preview)} of {total} "
                              f"rows. Export the full frame to a workspace file (e.g. "
                              f"pandas_frames__export) and use the file instead.")
        return _rows_from_value(preview)
    if isinstance(value, dict):
        columns = value.get("columns") or value.get("column_names")
        for key in ("rows", "data", "records", "results", "items", "values"):
            inner = value.get(key)
            if isinstance(inner, list):
                if inner and isinstance(inner[0], (list, tuple)) and isinstance(columns, list):
                    names = [c if isinstance(c, str) else (c.get("name") or str(c))
                             for c in columns]
                    return [dict(zip(names, row)) for row in inner][:MAX_ROWS]
                value = inner
                break
        else:
            # One object with scalar values is a one-row result.
            if value and all(not isinstance(v, (list, dict)) for v in value.values()):
                return [value]
            return None
    if isinstance(value, list):
        if not value:
            return []
        if all(isinstance(r, dict) for r in value):
            return value[:MAX_ROWS]
        if all(not isinstance(r, (list, dict)) for r in value):
            return [{"value": r} for r in value][:MAX_ROWS]
    return None


def _rows_from_markdown(text: str) -> list[dict] | None:
    lines = [l.strip() for l in text.splitlines() if l.strip().startswith("|")]
    if len(lines) < 2 or not re.match(r"^\|?\s*:?-{2,}", lines[1].replace("|", "", 1)):
        return None
    split = lambda l: [c.strip() for c in l.strip().strip("|").split("|")]
    header = split(lines[0])
    return [dict(zip(header, (_coerce(c) for c in split(l)))) for l in lines[2:]][:MAX_ROWS]


def _rows_from_csv(text: str) -> list[dict] | None:
    lines = text.splitlines()
    if len(lines) < 2:
        return None
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=",;\t|")
    except csv.Error:
        return None
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames or len(reader.fieldnames) < 2:
        return None
    return [{k: _coerce(v) for k, v in row.items() if k is not None}
            for _, row in zip(range(MAX_ROWS), reader)]


def _coerce(value: Any) -> Any:
    """'1 250,00' stays text; '1250.00' becomes a number — a chart needs to know which."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if re.fullmatch(r"-?\d+", text):
        try:
            return int(text)
        except ValueError:
            return text
    if re.fullmatch(r"-?\d*\.\d+(e-?\d+)?|-?\d+e-?\d+", text, re.IGNORECASE):
        try:
            return float(text)
        except ValueError:
            return text
    if text.lower() in ("null", "none", ""):
        return None
    return text


def rows_from_file(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).active
        values = sheet.iter_rows(values_only=True)
        header = [str(h) if h is not None else f"col{i + 1}" for i, h in enumerate(next(values, []))]
        return [dict(zip(header, row)) for _, row in zip(range(MAX_ROWS), values)]
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    rows = rows_from_text(text)
    if rows is None:
        raise SourceError(f"{path.name} does not hold a table this can read "
                          f"(expected CSV, JSON rows or XLSX).")
    return rows


def columns_of(rows: list[dict]) -> list[str]:
    seen: dict[str, None] = {}
    for row in rows[:500]:
        for key in row:
            seen.setdefault(str(key), None)
    return list(seen)


# ------------------------------------------------------------------ resolving

def resolve(source: Any, *, blocks: list[dict], history: list[list[dict]] | None = None,
            workspace: Path, charts=None, conversation_id: str = "") -> tuple[list[dict], str]:
    """Rows for a data reference, and a label saying where they came from.

    `#4` is looked up in this run first and then in earlier answers of the conversation,
    newest first — "export that table" in a follow-up means the table the reader saw.
    """
    if isinstance(source, dict):
        # Vega-Lite's own shape, {"name": "#7", "format": …}, or {"ref": "#7"}: the reference
        # is inside; the rest describes a fetch this app does not make.
        inner = next((source[k] for k in ("name", "ref", "source", "data", "url")
                      if isinstance(source.get(k), str) and source[k].strip()), None)
        if isinstance(source.get("values"), list):
            source = source["values"]
        elif inner is not None:
            source = inner
    if isinstance(source, list):
        rows = _rows_from_value(source)
        if rows is None:
            raise SourceError("Inline data must be a list of objects, one per row.")
        return rows, "inline values"
    text = str(source or "").strip()
    if not text:
        raise SourceError("No data source given. Name one: '#4' for the result of call #4, "
                          "'chart:c1' for a chart's data, or a workspace file path.")
    match = _REF.match(text)
    if match:
        ref = f"#{match.group(1)}"
        for scope in [blocks, *(history or [])]:
            for block in reversed(scope):
                # Never the call still running — that is the one asking for the rows.
                if block.get("type") == "tool" and block.get("ref") == ref and block.get("status") != "running":
                    return _rows_of_block(block, ref)
        known = [b["ref"] for b in blocks if b.get("type") == "tool" and b.get("ref")]
        raise SourceError(f"There is no call {ref} in this conversation. Calls with results "
                          f"right now: {', '.join(known[-12:]) or 'none'}.")
    if text.lower().startswith("chart:"):
        if charts is None:
            raise SourceError("Charts are not available here.")
        chart = charts.latest(conversation_id, text.split(":", 1)[1].strip())
        if not chart:
            raise SourceError(f"No chart named {text.split(':', 1)[1].strip()} in this conversation.")
        return list(chart["data"]), f"data of chart {chart['id']}"
    path = text.removeprefix("workspace:").strip()
    from app.tools import files as file_tool
    try:
        target = file_tool.resolve(workspace, path)
    except file_tool.OutsideWorkspace as exc:
        raise SourceError(str(exc)) from exc
    if not target.is_file():
        raise SourceError(f"'{path}' is not a file in the workspace.")
    return rows_from_file(target), path


def _rows_of_block(block: dict, ref: str) -> tuple[list[dict], str]:
    if not block.get("ok"):
        raise SourceError(f"Call {ref} ({block.get('name')}) failed, so it has no rows to use.")
    label = f"{ref} {block.get('name', '')}".strip()
    data = block.get("data")
    if isinstance(data, (list, dict)):
        rows = _rows_from_value(data)
        if rows:
            return rows, label
    # A result too large for the context was parked on disk whole; read that, not the
    # excerpt the model saw.
    if block.get("offloaded"):
        from app.deps import container
        path = container.workspace() / block["offloaded"]
        if path.is_file():
            rows = rows_from_text(path.read_text(encoding="utf-8", errors="replace"))
            if rows is not None:
                return rows, label
    rows = rows_from_text(block.get("text") or "")
    if rows is None:
        hint = (" Run it again ending with print(df.to_json(orient='records')) and use that "
                "call's #ref." if block.get("name") == "run_python" else "")
        raise SourceError(f"Call {ref} ({block.get('name')}) did not return a table. "
                          f"It starts: {(block.get('text') or '')[:120]!r}.{hint}")
    return rows, label
